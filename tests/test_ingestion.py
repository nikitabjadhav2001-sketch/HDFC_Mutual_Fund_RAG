import hashlib
import io
import json
import uuid
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import delete, func, select

from app.config import get_settings
from app.core.embedding import (
    EmbeddingError,
    HuggingFaceEmbeddingProvider,
    LocalEmbeddingProvider,
    OpenAIEmbeddingProvider,
    build_embedder,
)
from app.db import SessionLocal
from app.ingestion.parsers import ParseError, detect_format, parse
from app.ingestion.pipeline import ingest
from app.main import app
from app.models.tables import EMBEDDING_DIM, Document
from app.models.tables import Chunk as ChunkRow

client = TestClient(app)


def _db_ok() -> bool:
    from sqlalchemy import create_engine, text

    try:
        engine = create_engine(
            get_settings().database_url, connect_args={"connect_timeout": 2}
        )
        with engine.connect() as conn:
            conn.execute(text("SELECT 1"))
        engine.dispose()
        return True
    except Exception:
        return False


requires_db = pytest.mark.skipif(not _db_ok(), reason="PostgreSQL not reachable")


class FakeEmbedder:
    model = "fake-embedding"

    def __init__(self) -> None:
        self.calls: list[list[str]] = []

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        return [[0.1] * EMBEDDING_DIM for _ in texts]


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            "/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_docx() -> bytes:
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading("Intro", level=1)
    document.add_paragraph("Intro body text.")
    document.add_heading("Details", level=2)
    document.add_paragraph("Details body text.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- parsers -----------------------------------------------------------------


def test_parse_txt_uses_filename_as_title() -> None:
    doc = parse(b"plain text body", "txt", filename="notes.txt")
    assert doc.title == "notes"
    assert len(doc.sections) == 1
    assert doc.sections[0].text == "plain text body"


def test_parse_md_splits_headings_and_takes_h1_as_title() -> None:
    payload = "# Doc Title\n\nintro text\n\n## Section A\n\nalpha beta\n"
    doc = parse(payload.encode(), "md", filename="ignored.md")
    assert doc.title == "Doc Title"
    assert [(s.heading, s.text) for s in doc.sections] == [
        (None, "intro text"),
        ("Section A", "alpha beta"),
    ]


def test_parse_pdf_extracts_text_with_page_numbers() -> None:
    doc = parse(minimal_pdf("Hello Phase One"), "pdf", filename="report.pdf")
    assert "Hello Phase One" in doc.sections[0].text
    assert doc.sections[0].page == 1


def test_parse_docx_builds_heading_paths() -> None:
    doc = parse(make_docx(), "docx", filename="guide.docx")
    assert doc.title == "guide"
    assert [s.heading for s in doc.sections] == ["Intro", "Intro > Details"]
    assert "Intro body text." in doc.sections[0].text
    assert "Details body text." in doc.sections[1].text


def test_parse_invalid_pdf_raises_parse_error() -> None:
    with pytest.raises(ParseError):
        parse(b"this is not a pdf", "pdf", filename="broken.pdf")


def test_parse_empty_file_raises() -> None:
    with pytest.raises(ParseError):
        parse(b"", "txt", filename="empty.txt")


def test_parse_unsupported_format_raises() -> None:
    with pytest.raises(ParseError):
        parse(b"data", "exe", filename="app.exe")


def test_detect_format_normalizes_extensions() -> None:
    assert detect_format("a.markdown") == "md"
    assert detect_format("a.MD") == "md"
    assert detect_format("a.pdf") == "pdf"
    with pytest.raises(ParseError):
        detect_format("a.bin")


# --- embeddings --------------------------------------------------------------


def _embedding_provider(handler, **kwargs) -> OpenAIEmbeddingProvider:
    client = httpx.Client(transport=httpx.MockTransport(handler), base_url="https://test")
    return OpenAIEmbeddingProvider(
        "test-model", "test-key", backoff=0, client=client, **kwargs
    )


def _embedding_response(texts: list[str]) -> httpx.Response:
    order = list(range(len(texts)))[::-1]  # deliberately out of order
    data = [
        {"index": index, "embedding": [float(len(texts[index]))]} for index in order
    ]
    return httpx.Response(
        200, json={"data": data, "usage": {"total_tokens": len(texts)}}
    )


def test_embedding_batches_and_preserves_order() -> None:
    seen: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        texts = list((json.loads(request.content)["input"]))
        seen.append(len(texts))
        return _embedding_response(texts)

    provider = _embedding_provider(handler, batch_size=100)
    texts = ["x" * (i + 1) for i in range(250)]
    vectors = provider.embed_texts(texts)
    assert seen == [100, 100, 50]
    assert len(vectors) == len(texts)
    assert [v[0] for v in vectors] == [float(len(t)) for t in texts]


def test_embedding_retries_on_rate_limit() -> None:
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if calls == 1:
            return httpx.Response(429, text="slow down")
        return _embedding_response(list((json.loads(request.content)["input"])))

    provider = _embedding_provider(handler, max_retries=3)
    vectors = provider.embed_texts(["hello"])
    assert calls == 2
    assert vectors == [[5.0]]


def test_embedding_does_not_retry_client_errors() -> None:
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return httpx.Response(400, text="bad request")

    provider = _embedding_provider(handler, max_retries=3)
    with pytest.raises(EmbeddingError):
        provider.embed_texts(["hello"])
    assert calls == 1


def test_embedding_raises_after_exhausting_retries() -> None:
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return httpx.Response(500, text="boom")

    provider = _embedding_provider(handler, max_retries=1)
    with pytest.raises(EmbeddingError):
        provider.embed_texts(["hello"])
    assert calls == 2


def test_embedding_empty_input_makes_no_requests() -> None:
    def handler(request: httpx.Request) -> httpx.Response:  # pragma: no cover
        raise AssertionError("should not be called")

    provider = _embedding_provider(handler)
    assert provider.embed_texts([]) == []


# --- huggingface embeddings --------------------------------------------------


def _hf_provider(handler, token: str = "test-token", **kwargs) -> HuggingFaceEmbeddingProvider:
    client = httpx.Client(transport=httpx.MockTransport(handler), base_url="https://test")
    return HuggingFaceEmbeddingProvider(
        "sentence-transformers/all-MiniLM-L6-v2",
        token,
        backoff=0,
        client=client,
        **kwargs,
    )


def test_hf_embedding_batches_payload_and_order() -> None:
    seen: list[int] = []

    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == (
            "/models/sentence-transformers/all-MiniLM-L6-v2/pipeline/feature-extraction"
        )
        assert request.headers["Authorization"] == "Bearer test-token"
        body = json.loads(request.content)
        assert body["normalize"] is True
        assert body["truncate"] is True
        texts = list(body["inputs"])
        seen.append(len(texts))
        return httpx.Response(200, json=[[float(len(text))] for text in texts])

    provider = _hf_provider(handler, batch_size=2)
    texts = ["a", "bb", "ccc", "dddd"]
    vectors = provider.embed_texts(texts)
    assert seen == [2, 2]
    assert [v[0] for v in vectors] == [1.0, 2.0, 3.0, 4.0]


def test_hf_embedding_omits_auth_header_without_token() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert "Authorization" not in request.headers
        texts = json.loads(request.content)["inputs"]
        return httpx.Response(200, json=[[0.1] for _ in texts])

    provider = _hf_provider(handler, token="")
    assert provider.embed_texts(["hello"]) == [[0.1]]


def test_hf_embedding_tolerates_flattened_single_response() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert len(json.loads(request.content)["inputs"]) == 1
        return httpx.Response(200, json=[1.0, 2.0, 3.0])

    provider = _hf_provider(handler)
    assert provider.embed_texts(["hello"]) == [[1.0, 2.0, 3.0]]


def test_hf_embedding_retries_while_model_loads() -> None:
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        if calls == 1:
            return httpx.Response(503, json={"error": "Model is loading", "estimated_time": 1.0})
        texts = json.loads(request.content)["inputs"]
        return httpx.Response(200, json=[[0.5] for _ in texts])

    provider = _hf_provider(handler, max_retries=3)
    vectors = provider.embed_texts(["hello"])
    assert calls == 2
    assert vectors == [[0.5]]


def test_hf_embedding_does_not_retry_auth_errors() -> None:
    calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal calls
        calls += 1
        return httpx.Response(401, text="invalid token")

    provider = _hf_provider(handler, max_retries=3)
    with pytest.raises(EmbeddingError):
        provider.embed_texts(["hello"])
    assert calls == 1


def test_hf_embedding_rejects_vector_count_mismatch() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=[[1.0, 2.0]])

    provider = _hf_provider(handler)
    with pytest.raises(EmbeddingError, match="1 vectors for 2 inputs"):
        provider.embed_texts(["a", "b"])


def test_hf_embedding_rejects_inconsistent_dims() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json=[[1.0, 2.0], [3.0]])

    provider = _hf_provider(handler)
    with pytest.raises(EmbeddingError, match="inconsistent"):
        provider.embed_texts(["a", "b"])


def test_hf_embedding_empty_input_makes_no_requests() -> None:
    def handler(request: httpx.Request) -> httpx.Response:  # pragma: no cover
        raise AssertionError("should not be called")

    provider = _hf_provider(handler)
    assert provider.embed_texts([]) == []


def test_build_embedder_uses_huggingface(monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "embedding_provider", "huggingface")
    monkeypatch.setattr(settings, "embedding_model", "sentence-transformers/all-MiniLM-L6-v2")
    monkeypatch.setattr(settings, "hf_token", "hf_secret")
    monkeypatch.setattr(settings, "hf_base_url", "https://router.huggingface.co/hf-inference")

    embedder = build_embedder()
    assert isinstance(embedder, HuggingFaceEmbeddingProvider)
    assert embedder.model == "sentence-transformers/all-MiniLM-L6-v2"
    assert embedder.token == "hf_secret"


# --- local (sentence-transformers) embeddings --------------------------------


class _FakeEncoder:
    def __init__(self) -> None:
        self.calls: list[tuple[list[str], dict]] = []

    def encode(self, texts: list[str], **kwargs):
        self.calls.append((list(texts), kwargs))
        return [[float(len(text)), 1.0] for text in texts]


class _BoomEncoder:
    def encode(self, *args, **kwargs):  # pragma: no cover
        raise AssertionError("encode must not be called")


def test_local_embedding_converts_vectors_to_float_lists() -> None:
    encoder = _FakeEncoder()
    provider = LocalEmbeddingProvider(encoder=encoder)

    vectors = provider.embed_texts(["a", "bb"])

    assert vectors == [[1.0, 1.0], [2.0, 1.0]]
    texts, kwargs = encoder.calls[0]
    assert texts == ["a", "bb"]
    assert kwargs["normalize_embeddings"] is True
    assert kwargs["show_progress_bar"] is False


def test_local_embedding_empty_input_skips_encoder() -> None:
    provider = LocalEmbeddingProvider(encoder=_BoomEncoder())
    assert provider.embed_texts([]) == []


def test_build_embedder_uses_local(monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "embedding_provider", "local")
    monkeypatch.setattr(settings, "embedding_model", "sentence-transformers/all-MiniLM-L6-v2")

    embedder = build_embedder()
    assert isinstance(embedder, LocalEmbeddingProvider)
    assert embedder.model == "sentence-transformers/all-MiniLM-L6-v2"


# --- CLI arg handling (no DB needed) -----------------------------------------


def test_cli_rejects_missing_and_unsupported_files(tmp_path: Path) -> None:
    from app.ingest import main

    assert main([str(tmp_path / "nope.pdf")]) == 2
    weird = tmp_path / "app.exe"
    weird.write_bytes(b"data")
    assert main([str(weird)]) == 2


# --- DB-backed pipeline + API (require PostgreSQL) ---------------------------

LONG_MD = "\n\n".join(
    f"## Section {i}\n\n" + " ".join(f"word{j}" for j in range(60))
    for i in range(4)
)


@pytest.fixture()
def db():
    session = SessionLocal()
    session.execute(delete(ChunkRow))
    session.execute(delete(Document))
    session.commit()
    yield session
    session.execute(delete(ChunkRow))
    session.execute(delete(Document))
    session.commit()
    session.close()


def _add_document(db, data: bytes, filename: str, tmp_path: Path) -> uuid.UUID:
    fmt = detect_format(filename)
    source = tmp_path / filename
    source.write_bytes(data)
    doc = Document(
        id=uuid.uuid4(),
        title=Path(filename).stem,
        format=fmt,
        content_hash=hashlib.sha256(data).hexdigest(),
        status="queued",
        source_uri=str(source),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc.id


def _chunk_count(db) -> int:
    return db.execute(select(func.count()).select_from(ChunkRow)).scalar_one()


@requires_db
def test_ingest_markdown_reaches_done_with_chunks(db, tmp_path: Path) -> None:
    doc_id = _add_document(db, LONG_MD.encode(), "kb.md", tmp_path)
    embedder = FakeEmbedder()
    result = ingest(doc_id, db, embedder=embedder)
    assert result.status == "done"
    assert result.error is None
    assert _chunk_count(db) > 0
    assert len(embedder.calls) == 1


@requires_db
def test_reingest_is_idempotent(db, tmp_path: Path) -> None:
    doc_id = _add_document(db, LONG_MD.encode(), "kb.md", tmp_path)
    embedder = FakeEmbedder()
    ingest(doc_id, db, embedder=embedder)
    count_after_first = _chunk_count(db)

    result = ingest(doc_id, db, embedder=embedder)
    assert result.status == "done"
    assert _chunk_count(db) == count_after_first
    assert len(embedder.calls) == 1, "unchanged content must not be re-embedded"


@requires_db
def test_parse_failure_marks_document_failed(db, tmp_path: Path) -> None:
    doc_id = _add_document(db, b"not a real pdf", "broken.pdf", tmp_path)
    result = ingest(doc_id, db, embedder=FakeEmbedder())
    assert result.status == "failed"
    assert result.error
    assert _chunk_count(db) == 0


@requires_db
def test_missing_source_marks_document_failed(db, tmp_path: Path) -> None:
    doc_id = _add_document(db, LONG_MD.encode(), "kb.md", tmp_path)
    Path(tmp_path / "kb.md").unlink()
    result = ingest(doc_id, db, embedder=FakeEmbedder())
    assert result.status == "failed"
    assert "missing" in (result.error or "")


@requires_db
def test_admin_upload_status_delete_flow(db, tmp_path: Path, monkeypatch) -> None:
    async def _noop_enqueue(document_id: str) -> None:
        return None

    monkeypatch.setattr("app.api.routes_admin.enqueue_ingest", _noop_enqueue)
    monkeypatch.setattr(get_settings(), "storage_dir", str(tmp_path / "storage"))

    payload = LONG_MD.encode()

    response = client.post(
        "/api/v1/admin/documents", files={"file": ("kb.md", payload, "text/markdown")}
    )
    assert response.status_code == 201
    body = response.json()
    assert body["created"] is True
    document_id = body["document"]["id"]
    assert body["document"]["status"] == "queued"

    ingest(uuid.UUID(document_id), db, embedder=FakeEmbedder())

    status = client.get(f"/api/v1/admin/documents/{document_id}/status").json()
    assert status["status"] == "done"
    assert status["chunk_count"] > 0
    assert _chunk_count(db) == status["chunk_count"]

    duplicate = client.post(
        "/api/v1/admin/documents", files={"file": ("kb-copy.md", payload, "text/markdown")}
    )
    assert duplicate.status_code == 201
    assert duplicate.json()["created"] is False
    assert duplicate.json()["document"]["id"] == document_id

    listing = client.get("/api/v1/admin/documents").json()
    assert len(listing) == 1
    assert listing[0]["id"] == document_id

    deleted = client.delete(f"/api/v1/admin/documents/{document_id}")
    assert deleted.status_code == 204
    assert client.get("/api/v1/admin/documents").json() == []
    assert _chunk_count(db) == 0
    remaining = db.execute(
        select(Document).where(Document.id == uuid.UUID(document_id))
    ).scalar_one_or_none()
    assert remaining is None


@requires_db
def test_admin_delete_unknown_document_is_404(db) -> None:
    assert client.delete(f"/api/v1/admin/documents/{uuid.uuid4()}").status_code == 404
    assert client.get(f"/api/v1/admin/documents/{uuid.uuid4()}/status").status_code == 404


# --- API validation (no DB needed) -------------------------------------------


def test_upload_rejects_unsupported_format() -> None:
    response = client.post(
        "/api/v1/admin/documents",
        files={"file": ("app.exe", b"binary", "application/octet-stream")},
    )
    assert response.status_code == 415


def test_upload_rejects_empty_file() -> None:
    response = client.post(
        "/api/v1/admin/documents", files={"file": ("notes.md", b"", "text/markdown")}
    )
    assert response.status_code == 400
